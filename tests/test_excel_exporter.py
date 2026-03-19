"""Integration test for Excel exporter using synthetic data – no browser required."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

import pytest
import tempfile
import openpyxl
from datetime import datetime

from models.product import Product, SKUVariant
from models.supplier import Supplier, SupplierSKU, WangwangData, SourcingResult
from output.excel_exporter import ExcelExporter


def _build_result() -> SourcingResult:
    product = Product(
        source_url="https://test-shop.myshopify.com/products/widget",
        title="Test Widget",
        platform="shopify",
    )
    product.skus = [
        SKUVariant(sku_id="1", attributes={"Color": "Red"}, price=29.99),
        SKUVariant(sku_id="2", attributes={"Color": "Blue"}, price=27.99),
    ]

    def _supplier(rank, name, price, url):
        s = Supplier(
            shop_name=name,
            product_url=url,
            shop_url=f"https://shop.1688.com/{name}",
            rank_position=rank,
            rank_score=1.0 - rank * 0.1,
            rating=4.5,
            years_operating=5,
        )
        s.skus = [
            SupplierSKU(
                sku_id=f"s{rank}_1",
                attributes={"颜色": "红色"},
                price_cny=price,
                weight_net_kg=0.3,
                weight_gross_kg=0.4,
                package_size_cm="15×10×8cm",
                inventory=1000,
                ww_price_cny=price - 0.5,
                ww_weight_net_kg=0.31,
                ww_weight_gross_kg=0.41,
                ww_package_size_cm="15×10×8cm",
                ww_moq=50,
            )
        ]
        s.wangwang_id = f"ww_{name}"
        s.wangwang_data = WangwangData(
            inquiry_sent=True,
            inquiry_time=datetime(2024, 1, 1, 10, 0),
            reply_received=True,
            reply_time=datetime(2024, 1, 1, 10, 5),
            raw_reply=f"报价{price}元，净重0.3kg，MOQ50件",
            moq=50,
        )
        s.data_source = "both"
        return s

    sup1 = _supplier(1, "供应商A", 12.5, "https://detail.1688.com/offer/111.html")
    sup2 = _supplier(2, "供应商B", 13.5, "https://detail.1688.com/offer/222.html")
    sup3 = _supplier(3, "供应商C", 14.0, "https://detail.1688.com/offer/333.html")

    return SourcingResult(
        product=product,
        suppliers=[sup1, sup2, sup3],
        recommended_supplier=sup1,
    )


def test_excel_creates_file():
    result = _build_result()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = Path(f.name)
    exporter = ExcelExporter()
    output = exporter.export([result], out_path)
    assert output.exists()
    assert output.stat().st_size > 0


def test_excel_has_three_sheets():
    result = _build_result()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = Path(f.name)
    exporter = ExcelExporter()
    output = exporter.export([result], out_path)
    wb = openpyxl.load_workbook(str(output))
    assert len(wb.sheetnames) == 3


def test_excel_summary_has_data_row():
    result = _build_result()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = Path(f.name)
    exporter = ExcelExporter()
    output = exporter.export([result], out_path)
    wb = openpyxl.load_workbook(str(output))
    ws = wb.worksheets[0]
    # Row 1 = header, row 2 = first data row
    assert ws.max_row >= 2
    # Source URL should be in column A row 2
    assert "test-shop" in (ws["A2"].value or "")


def test_excel_sku_detail_rows():
    result = _build_result()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        out_path = Path(f.name)
    exporter = ExcelExporter()
    output = exporter.export([result], out_path)
    wb = openpyxl.load_workbook(str(output))
    ws = wb.worksheets[1]  # SKU Detail sheet
    # 3 suppliers × 1 SKU each = 3 data rows + 1 header
    assert ws.max_row >= 4
